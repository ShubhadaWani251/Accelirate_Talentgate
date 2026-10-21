"""The export returns the rows the list is showing, not every row in the database.

Reported by a tester: "Exporting data from Passed Candidates currently exports all
candidates." The export view honoured batch_id and the date range and nothing else, so the
Export button beside a filtered table quietly ignored the filter it sat next to.

Both endpoints now run the same views/candidates.apply_candidate_filters, so these assert the
two agree rather than asserting the export's filtering in isolation - a copy of the logic
passing its own tests is exactly the failure mode that produced the bug.
"""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from api.models import Candidate

pytestmark = pytest.mark.django_db


def _exported_emails(response):
    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    email_col = headers.index('Email')
    return {row[email_col].value for row in ws.iter_rows(min_row=2)}


def _listed_emails(response):
    return {row['email'] for row in response.data['results']}


@pytest.fixture
def cohort(ta_user, make_batch, make_candidate):
    """Three candidates in one batch, one per result band."""
    batch = make_batch(ta_user)
    return {
        result: make_candidate(batch, ta_user, result=result)
        for result in (Candidate.Result.PASS, Candidate.Result.FAIL,
                       Candidate.Result.BORDERLINE)
    }


class TestExportHonoursTheListFilters:
    def test_result_filter_reaches_the_export(self, ta_user, client_for, cohort):
        client = client_for(ta_user)

        response = client.get('/api/candidates/export/', {'result': Candidate.Result.PASS})

        assert response.status_code == 200
        assert _exported_emails(response) == {cohort[Candidate.Result.PASS].email}

    def test_export_matches_the_list_it_was_opened_from(self, ta_user, client_for, cohort):
        """The property that actually matters, and the one the shared seam exists to hold."""
        client = client_for(ta_user)
        params = {'result': Candidate.Result.PASS}

        listed = _listed_emails(client.get('/api/candidates/', params))
        exported = _exported_emails(client.get('/api/candidates/export/', params))

        assert exported == listed

    def test_name_filter_reaches_the_export(self, ta_user, make_batch, make_candidate,
                                            client_for):
        batch = make_batch(ta_user)
        wanted = make_candidate(batch, ta_user, first_name='Priya')
        make_candidate(batch, ta_user, first_name='Devesh')
        client = client_for(ta_user)

        response = client.get('/api/candidates/export/', {'name': 'Priya'})

        assert _exported_emails(response) == {wanted.email}

    def test_score_range_reaches_the_export(self, ta_user, make_batch, make_candidate,
                                            make_invitation, client_for):
        from api.models import ExamAttempt

        batch = make_batch(ta_user)
        high = make_candidate(batch, ta_user)
        low = make_candidate(batch, ta_user)
        for candidate, marks in ((high, 30), (low, 5)):
            ExamAttempt.objects.create(
                candidate=candidate, invitation=make_invitation(candidate, ta_user),
                status=ExamAttempt.Status.SUBMITTED, total_marks_earned=marks, total_marks=40,
            )
        client = client_for(ta_user)

        response = client.get('/api/candidates/export/', {'score_min': '20'})

        assert _exported_emails(response) == {high.email}

    def test_no_filters_still_exports_everyone(self, ta_user, client_for, cohort):
        """The Dashboard and sidebar Export actions pass nothing and must keep meaning 'all'."""
        client = client_for(ta_user)

        response = client.get('/api/candidates/export/')

        assert _exported_emails(response) == {c.email for c in cohort.values()}
